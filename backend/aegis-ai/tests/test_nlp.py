import asyncio
import time

from runtime.agents import AgentCapability, AgentDescriptor, AgentRequest, AgentResult, AgentStatus, BaseAgent, AgentRegistry, MasterAgent, OllamaSpecialistAgent
from runtime.nlp import NLPPreprocessor


def test_cleaning_and_original_preservation():
    raw = "Identify   the damaged   guard\nnear P-101.\x00"
    result = NLPPreprocessor().process(raw)
    assert result.original_text == raw
    assert result.normalized_text == "Identify the damaged guard near P-101."
    assert "P-101" in result.entities


def test_ocr_and_terminology_corrections_are_confident():
    result = NLPPreprocessor().process("anlyze this inspction report and find majr findings in P & ID; q.c.")
    assert "analyze" in result.normalized_text
    assert "inspection" in result.normalized_text
    assert "major" in result.normalized_text
    assert "P&ID" in result.normalized_text
    assert "quality control" in result.normalized_text
    assert all(c["confidence"] >= 0.9 for c in result.corrections)


def test_identifiers_and_ambiguous_tags_are_preserved():
    result = NLPPreprocessor().process("check P-101 and FT-101, but P-10I is unclear")
    assert "P-101" in result.normalized_text and "FT-101" in result.normalized_text
    assert "P-10I" in result.normalized_text
    ambiguous = [c for c in result.corrections if c["original"] == "P-10I"]
    assert ambiguous and ambiguous[0]["normalized"] == "P-10I"
    assert ambiguous[0]["confidence"] < 0.5


def test_prompt_enhancement_is_proportional():
    vague = NLPPreprocessor().process("what is this image")
    assert vague.metadata["enhancement_applied"] is True
    assert "provided image" in vague.enhanced_prompt
    direct = NLPPreprocessor().process("What is P&ID?")
    assert direct.enhanced_prompt == direct.normalized_text


def test_intent_entities_and_speed():
    started = time.perf_counter()
    result = NLPPreprocessor().process("Analyze this inspection PDF and find major safety issues.")
    assert (time.perf_counter() - started) < 0.1
    assert result.intent == "document_analysis"
    assert result.modality == "document"
    assert "major" in result.severity_terms


class _CaptureAgent(BaseAgent):
    def __init__(self):
        self.descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
        self.seen = None

    async def run(self, request):
        self.seen = request
        return AgentResult(agent="document_agent", status=AgentStatus.SUCCESS, summary="ok")


def test_master_passes_preprocessed_context_to_specialist():
    agent = _CaptureAgent()
    registry = AgentRegistry()
    registry.register(agent)

    async def planner(request, _caps):
        assert "inspection" in request
        return [{"agent": "document_agent", "task": request}]

    state = asyncio.run(MasterAgent(registry, planner=planner).run("anlyze this inspction report"))
    assert state.user_request == "anlyze this inspction report"
    assert state.preprocessing["normalized_text"] == "analyze this inspection report"
    assert agent.seen.context["nlp"]["original_text"] == state.user_request


def test_document_ocr_keeps_raw_and_adds_normalized_text():
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {"document_runner": lambda _: {"status": "complete", "ocr_text": "Severity: majr."}})
    result = asyncio.run(agent.run(AgentRequest(task="inspect", context={"input_path": "x.pdf"})))
    assert result.result["raw_ocr_text"] == "Severity: majr."
    assert result.result["normalized_ocr_text"] == "Severity: major."


def test_document_agent_creates_and_verifies_docx(tmp_path):
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {})
    result = asyncio.run(agent.run(AgentRequest(
        task="create a document of Mahatma Gandhi and save it in docx",
        context={"workspace_root": str(tmp_path)},
    )))
    assert result.status == AgentStatus.SUCCESS
    assert result.metadata["operation"] == "create_document"
    artifact = result.result["artifact"]["path"]
    assert artifact.endswith(".docx") and "Mahatma Gandhi" in result.summary
    assert result.verification == result.result["verification"]


def test_document_analysis_without_input_path_is_clear_failure():
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {"document_runner": lambda _: {"status": "complete"}})
    result = asyncio.run(agent.run(AgentRequest(task="analyze existing document", context={})))
    assert result.status == AgentStatus.FAILURE
    assert "missing input_path" in result.errors


def test_markdown_and_plain_text_docx_content_is_readable(tmp_path):
    from pipeline.docx_writer import write_text_document
    from docx import Document
    path = tmp_path / "content.docx"
    write_text_document("Agentic AI\n\n# Introduction\n\nAgents can plan.\n\n- tool use\n- verification\n\n# Conclusion\n\nDone.", path)
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert all(value in text for value in ("Agentic AI", "Introduction", "Agents can plan.", "tool use", "verification", "Conclusion"))
    plain = tmp_path / "plain.docx"
    write_text_document("This is a simple document.\nIt contains multiple paragraphs.", plain)
    assert "multiple paragraphs" in "\n".join(p.text for p in Document(str(plain)).paragraphs)


def test_empty_docx_content_fails_before_creation(tmp_path):
    from pipeline.docx_writer import write_text_document, DocxWriterError
    import pytest
    with pytest.raises(DocxWriterError, match="content_generation_failed"):
        write_text_document("", tmp_path / "empty.docx")


def test_document_creation_uses_local_model_content(tmp_path):
    from unittest.mock import AsyncMock, MagicMock
    from docx import Document
    provider = MagicMock()
    provider.generate = AsyncMock(return_value=MagicMock(content="Agentic AI\n\n# Introduction\n\nGenerated by the local model."))
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, provider, {})
    result = asyncio.run(agent.run(AgentRequest(task="create a document about Agentic AI and save it in docx", context={"workspace_root": str(tmp_path)})))
    assert result.status == AgentStatus.SUCCESS
    assert result.metadata["generation_source"] == "local_model"
    text = "\n".join(p.text for p in Document(result.result["artifact"]["path"]).paragraphs)
    assert "Generated by the local model." in text
    provider.generate.assert_awaited_once()


def test_document_agent_creates_pdf_markdown_and_txt(tmp_path):
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {})
    for fmt, suffix in (("pdf", ".pdf"), ("markdown", ".md"), ("md", ".md"), ("txt", ".txt")):
        result = asyncio.run(agent.run(AgentRequest(task=f"create {fmt} about Agentic AI", context={"workspace_root": str(tmp_path)})))
        assert result.status == AgentStatus.SUCCESS
        assert result.result["artifact"]["path"].endswith(suffix)
        assert result.verification["status"] == "passed"


def test_document_topic_and_requirements_propagate_for_phrase(tmp_path):
    from pathlib import Path
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {})
    result = asyncio.run(agent.run(AgentRequest(
        task="create a txt file for car racing explain what is it and how it is planned",
        context={"workspace_root": str(tmp_path)},
    )))
    assert result.status == AgentStatus.SUCCESS
    assert result.metadata["topic"] == "car racing"
    text = Path(result.result["artifact"]["path"]).read_text()
    assert "car racing" in text.lower()
    assert "explain what is it and how it is planned" in result.metadata["content_requirements"].lower()


def test_document_agent_consumes_master_task_spec(tmp_path):
    descriptor = AgentDescriptor(name="document_agent", role="document", capabilities=[AgentCapability.DOCUMENT], provider_name="fake")
    agent = OllamaSpecialistAgent(descriptor, None, {})
    result = asyncio.run(agent.run(AgentRequest(
        task="create txt file", context={"workspace_root": str(tmp_path),
        "task_spec": {"operation": "create", "topic": "solar energy", "content_requirements": ["explain benefits"], "artifact_format": "txt"}},
    )))
    assert result.status == AgentStatus.SUCCESS
    assert result.metadata["topic"] == "solar energy"
    assert "solar energy" in open(result.result["artifact"]["path"], encoding="utf-8").read().lower()


def test_registry_advertises_document_creation_operation():
    from models.registry import ModelRegistry
    from runtime.agents import build_default_agent_registry
    registry = ModelRegistry.from_yaml("config/models.yaml")
    agents = build_default_agent_registry(registry, {})
    descriptor = next(d for d in agents.list() if d.name == "document_agent")
    assert "create_document" in descriptor.allowed_tools
