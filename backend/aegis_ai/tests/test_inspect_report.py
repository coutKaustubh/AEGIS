"""Comprehensive tests for the inspect-report pipeline.

Tests cover:
1. Input validation and routing
2. OCR with page markers
3. Strict findings schema validation
4. DOCX generation and content verification
5. Sandbox execution (success, timeout, unavailable)
6. Network policy (deny non-loopback, allow loopback)
7. Run metadata and artifact persistence
8. CLI command/help behavior
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from PIL import Image

# Ensure project root on path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from pipeline.network_policy import NetworkPolicy, _is_loopback
from pipeline.docx_writer import write_approval_note, HAS_DOCX, DocxWriterError
from pipeline.sandbox import run_sandboxed, SandboxUnavailableError, SandboxResult
from pipeline.inspect_report import (
    _validate_input,
    _extract_findings_deterministic,
    _detect_severity,
    run_inspect_report,
)
from runtime.errors import WorkbenchError
from tools.ocr.models import OCRTextBlock
from tools.ocr.paddle import PaddleOCRBackend


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

class FakeOCREngine:
    """Deterministic fake OCR engine for testing."""
    def predict(self, source: str):
        return [{
            "rec_texts": [
                "INSPECTION REPORT",
                "Date: 15/03/2025",
                "Site: Plant Alpha",
                "Inspector: John Doe",
                "Finding 1: Critical crack detected in vessel V-101",
                "Severity: Critical",
                "Corrective action required: Immediate shutdown and repair",
                "Finding 2: Minor corrosion on pipeline P-205",
                "Observation: Surface rust, monitoring recommended",
            ],
            "rec_scores": [0.98, 0.95, 0.92, 0.88, 0.91, 0.96, 0.90, 0.85, 0.80],
            "rec_boxes": [
                [10, 10, 200, 30],
                [10, 40, 200, 60],
                [10, 70, 200, 90],
                [10, 100, 200, 120],
                [10, 130, 400, 150],
                [10, 160, 200, 180],
                [10, 190, 400, 210],
                [10, 220, 400, 240],
                [10, 250, 400, 270],
            ],
        }]


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    """Create a minimal test image."""
    path = tmp_path / "test_report.png"
    Image.new("RGB", (100, 100), "white").save(path)
    return path


@pytest.fixture
def sample_jpg(tmp_path: Path) -> Path:
    path = tmp_path / "test_report.jpg"
    Image.new("RGB", (100, 100), "white").save(path)
    return path


@pytest.fixture
def fake_ocr_backend() -> PaddleOCRBackend:
    return PaddleOCRBackend(engine_factory=lambda **_: FakeOCREngine())


@pytest.fixture
def sample_findings() -> dict:
    return {
        "report_metadata": {
            "report_title": "Test Inspection Report",
            "inspection_date": "15/03/2025",
            "site": "Plant Alpha",
            "unit": "Unit 1",
            "inspector": "John Doe",
        },
        "findings": [
            {
                "id": "F-001",
                "description": "Critical crack detected in vessel V-101",
                "location_or_asset": "V-101",
                "severity": "critical",
                "evidence": "Visible crack along weld seam",
                "page": 1,
                "confidence": 0.91,
            },
            {
                "id": "F-002",
                "description": "Minor corrosion on pipeline P-205",
                "location_or_asset": "P-205",
                "severity": "minor",
                "evidence": "Surface rust observed",
                "page": 1,
                "confidence": 0.85,
            },
        ],
        "nonconformities": ["Weld quality below standard on V-101"],
        "corrective_actions": ["Immediate shutdown and repair of V-101"],
        "missing_information": [],
        "uncertainties": ["OCR confidence on some regions below 0.9"],
        "approval_recommendation": "Manual review required — automated extraction only",
    }


@pytest.fixture
def sandbox_script(tmp_path: Path) -> Path:
    script = tmp_path / "test_task.py"
    script.write_text("print('Hello from sandbox')\nprint(2 + 2)\n")
    return script


@pytest.fixture
def timeout_script(tmp_path: Path) -> Path:
    script = tmp_path / "slow_task.py"
    script.write_text("import time\ntime.sleep(60)\n")
    return script


# ---------------------------------------------------------------------------
# 1. Input validation and routing
# ---------------------------------------------------------------------------

class TestInputValidation:

    def test_valid_image_accepted(self, sample_image: Path) -> None:
        result = _validate_input(sample_image)
        assert result == sample_image.resolve()

    def test_valid_jpg_accepted(self, sample_jpg: Path) -> None:
        result = _validate_input(sample_jpg)
        assert result.suffix == ".jpg"

    def test_missing_file_rejected(self, tmp_path: Path) -> None:
        with pytest.raises(WorkbenchError, match="not found"):
            _validate_input(tmp_path / "nonexistent.pdf")

    def test_unsupported_extension_rejected(self, tmp_path: Path) -> None:
        f = tmp_path / "report.doc"
        f.write_text("test")
        with pytest.raises(WorkbenchError, match="Unsupported"):
            _validate_input(f)

    def test_pdf_extension_accepted(self, tmp_path: Path) -> None:
        f = tmp_path / "report.pdf"
        f.write_bytes(b"%PDF-1.4 test content")
        result = _validate_input(f)
        assert result.suffix == ".pdf"


# ---------------------------------------------------------------------------
# 3. OCR output with page markers
# ---------------------------------------------------------------------------

class TestOCRPageMarkers:

    def test_ocr_produces_page_markers(self, sample_image: Path, fake_ocr_backend: PaddleOCRBackend) -> None:
        from pipeline.inspect_report import _ocr_single_image
        text, blocks = _ocr_single_image(sample_image, fake_ocr_backend)
        assert "--- PAGE 1 ---" in text
        assert all(b.page_number == 1 for b in blocks)


# ---------------------------------------------------------------------------
# 4. Findings schema validation
# ---------------------------------------------------------------------------

class TestFindingsExtraction:

    @staticmethod
    def _extract(text: str) -> dict:
        return _extract_findings_deterministic(
            f"--- PAGE 1 ---\n{text}", [], Path("test.pdf")
        )

    def test_deterministic_extraction_returns_schema(self, sample_image: Path, fake_ocr_backend: PaddleOCRBackend) -> None:
        from pipeline.inspect_report import _ocr_single_image
        text, blocks = _ocr_single_image(sample_image, fake_ocr_backend)
        findings = _extract_findings_deterministic(text, blocks, sample_image)

        # Schema keys present
        assert "report_metadata" in findings
        assert "findings" in findings
        assert "uncertainties" in findings
        assert "nonconformities" in findings
        assert "corrective_actions" in findings
        assert "missing_information" in findings
        assert "approval_recommendation" in findings

    def test_findings_have_required_fields(self, sample_image: Path, fake_ocr_backend: PaddleOCRBackend) -> None:
        from pipeline.inspect_report import _ocr_single_image
        text, blocks = _ocr_single_image(sample_image, fake_ocr_backend)
        findings = _extract_findings_deterministic(text, blocks, sample_image)

        for f in findings["findings"]:
            assert "id" in f
            assert "description" in f
            assert "page" in f
            assert "evidence" in f
            assert "confidence" in f
            assert isinstance(f["page"], int)
            assert 0 <= f["confidence"] <= 1

    def test_severity_detection(self) -> None:
        assert _detect_severity("Critical failure detected") == "critical"
        assert _detect_severity("Major defect found") == "major"
        assert _detect_severity("Minor observation noted") == "minor"
        assert _detect_severity("Normal text") is None

    def test_explicit_minor_severity(self) -> None:
        result = self._extract("2. Defect: Gauge fluctuation. Severity: minor.")
        assert result["findings"][0]["severity"] == "minor"

    def test_explicit_major_severity(self) -> None:
        result = self._extract("2. Defect: Gauge fluctuation. Severity - major.")
        assert result["findings"][0]["severity"] == "major"

    def test_explicit_critical_severity(self) -> None:
        result = self._extract("2. Defect: Gauge fluctuation. severity=critical.")
        assert result["findings"][0]["severity"] == "critical"

    def test_next_line_critical_severity(self) -> None:
        result = self._extract(
            "3. Hazard: Missing safety guard.\nSeverity: critical. Immediate action required."
        )
        assert result["findings"][0]["severity"] == "critical"

    def test_nonconformity_is_finding(self) -> None:
        result = self._extract("1. Nonconformity: Corrosion on valve V-12.")
        assert result["findings"][0]["description"].startswith("1. Nonconformity")
        assert result["nonconformities"] == ["1. Nonconformity: Corrosion on valve V-12."]

    def test_nonconformance_is_finding(self) -> None:
        result = self._extract("1. Nonconformance: Corrosion on valve V-12.")
        assert result["findings"][0]["description"].startswith("1. Nonconformance")
        assert result["nonconformities"] == ["1. Nonconformance: Corrosion on valve V-12."]

    def test_ocr_prefixed_observation_is_finding(self) -> None:
        result = self._extract(". Observation: Guard missing.")
        assert result["findings"][0]["evidence"] == ". Observation: Guard missing."

    def test_recommendation_not_corrective_action(self) -> None:
        result = self._extract(
            "Recommendation: Inspect monthly.\nCorrective action: Install guard.\n"
            "Action required: Install sign.\nMust be: Completed immediately."
        )
        assert "Recommendation: Inspect monthly." not in result["corrective_actions"]
        assert result["corrective_actions"] == [
            "Corrective action: Install guard.",
            "Action required: Install sign.",
            "Must be: Completed immediately.",
        ]

    def test_no_invention_on_empty_text(self) -> None:
        findings = _extract_findings_deterministic("--- PAGE 1 ---\n", [], Path("test.pdf"))
        # Should not invent findings
        meta = findings["report_metadata"]
        assert meta["inspection_date"] is None or meta["inspection_date"] == "Unknown Report" or meta["inspection_date"] is None
        assert len(findings["uncertainties"]) > 0  # Should note uncertainties


# ---------------------------------------------------------------------------
# 5. DOCX creation
# ---------------------------------------------------------------------------

class TestDocxWriter:

    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    def test_docx_created_with_required_sections(self, tmp_path: Path, sample_findings: dict) -> None:
        output = tmp_path / "approval_note.docx"
        result = write_approval_note(sample_findings, output)
        assert result.exists()
        assert result.stat().st_size > 0

    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    def test_docx_contains_draft_marker(self, tmp_path: Path, sample_findings: dict) -> None:
        from docx import Document
        output = tmp_path / "approval_note.docx"
        write_approval_note(sample_findings, output)
        doc = Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "DRAFT" in full_text
        assert "HUMAN APPROVAL REQUIRED" in full_text

    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    def test_docx_has_findings_table(self, tmp_path: Path, sample_findings: dict) -> None:
        from docx import Document
        output = tmp_path / "approval_note.docx"
        write_approval_note(sample_findings, output)
        doc = Document(str(output))
        # Should have at least 3 tables (metadata, findings, signature)
        assert len(doc.tables) >= 3

    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    def test_docx_has_human_approval_section(self, tmp_path: Path, sample_findings: dict) -> None:
        from docx import Document
        output = tmp_path / "approval_note.docx"
        write_approval_note(sample_findings, output)
        doc = Document(str(output))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert any("Human Approval" in h for h in headings)


# ---------------------------------------------------------------------------
# 6. Sandbox tests
# ---------------------------------------------------------------------------

class TestSandbox:

    def test_successful_execution(self, sandbox_script: Path) -> None:
        result = run_sandboxed(script_path=sandbox_script, timeout_seconds=10)
        assert result.success is True
        assert result.exit_code == 0
        assert "Hello from sandbox" in result.stdout
        assert "4" in result.stdout
        assert result.timed_out is False
        assert result.sandbox_backend == "subprocess"

    def test_timeout_execution(self, timeout_script: Path) -> None:
        result = run_sandboxed(script_path=timeout_script, timeout_seconds=2)
        assert result.success is False
        assert result.timed_out is True
        assert result.sandbox_backend == "subprocess"

    def test_nonexistent_script_fails(self, tmp_path: Path) -> None:
        with pytest.raises(SandboxUnavailableError, match="not found"):
            run_sandboxed(script_path=tmp_path / "missing.py")

    def test_non_python_script_rejected(self, tmp_path: Path) -> None:
        f = tmp_path / "script.sh"
        f.write_text("echo hi")
        with pytest.raises(SandboxUnavailableError, match=".py"):
            run_sandboxed(script_path=f)

    def test_sandbox_result_serializable(self, sandbox_script: Path) -> None:
        result = run_sandboxed(script_path=sandbox_script, timeout_seconds=10)
        d = result.to_dict()
        assert json.dumps(d)  # Must be JSON-serializable


# ---------------------------------------------------------------------------
# 7. Network policy tests
# ---------------------------------------------------------------------------

class TestNetworkPolicy:

    def test_loopback_ollama_allowed(self) -> None:
        policy = NetworkPolicy(run_id="test-1", allowed_loopback_port=11434)
        assert policy.check_connection("127.0.0.1", 11434) is True

    def test_loopback_other_port_allowed(self) -> None:
        policy = NetworkPolicy(run_id="test-2", allowed_loopback_port=11434)
        assert policy.check_connection("127.0.0.1", 8080) is True

    def test_external_denied(self) -> None:
        policy = NetworkPolicy(run_id="test-3")
        assert policy.check_connection("8.8.8.8", 443) is False
        assert policy.denied_count == 1

    def test_url_check_external_denied(self) -> None:
        policy = NetworkPolicy(run_id="test-4")
        assert policy.check_url("https://api.openai.com/v1/chat") is False
        assert policy.denied_count == 1

    def test_url_check_localhost_allowed(self) -> None:
        policy = NetworkPolicy(run_id="test-5", allowed_loopback_port=11434)
        assert policy.check_url("http://localhost:11434/api/generate") is True

    def test_network_report_saved(self, tmp_path: Path) -> None:
        policy = NetworkPolicy(run_id="test-6")
        policy.check_connection("8.8.8.8", 443)
        policy.check_connection("127.0.0.1", 11434)
        path = policy.save(tmp_path)
        assert path.exists()

        report = json.loads(path.read_text())
        assert report["run_id"] == "test-6"
        assert report["policy_mode"] == "enforce_local_only"
        assert len(report["denied_attempts"]) == 1
        assert len(report["allowed_connections"]) == 1
        assert report["external_connection_count"] >= 1

    def test_summary_line_clean(self) -> None:
        policy = NetworkPolicy(run_id="test-7")
        assert policy.summary_line() == "EXTERNAL NETWORK CALLS: 0"

    def test_summary_line_with_denials(self) -> None:
        policy = NetworkPolicy(run_id="test-8")
        policy.check_connection("8.8.8.8", 443)
        assert "FAIL" in policy.summary_line()

    def test_ipv6_loopback(self) -> None:
        assert _is_loopback("::1") is True
        policy = NetworkPolicy(run_id="test-9")
        assert policy.check_connection("::1", 11434) is True


# ---------------------------------------------------------------------------
# 8. Full pipeline integration (with fake OCR)
# ---------------------------------------------------------------------------

class TestFullPipeline:

    @staticmethod
    def _fake_pdf_pages(pdf_path: Path, output_dir: Path) -> list[Path]:
        pages_dir = output_dir / "pages"
        pages_dir.mkdir()
        pages = [pages_dir / "page_0001.png", pages_dir / "page_0002.png"]
        for page in pages:
            page.write_bytes(b"fake page")
        return pages

    def test_pdf_page_images_are_cleaned_after_ocr(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        source_pdf = tmp_path / "inspection.pdf"
        source_pdf.write_bytes(b"%PDF-1.4")
        monkeypatch.setattr("pipeline.inspect_report._pdf_to_images", self._fake_pdf_pages)
        monkeypatch.setattr(
            "pipeline.inspect_report._ocr_images",
            lambda images, backend: ("--- PAGE 1 ---\nFinding: test", []),
        )

        out_dir = tmp_path / "run"
        result = run_inspect_report(input_path=source_pdf, output_dir=out_dir, skip_docx=True)

        assert result["status"] == "complete"
        assert source_pdf.exists()
        assert not (out_dir / "pages").exists()
        trace = json.loads((out_dir / "trace.json").read_text())
        assert any(phase["phase"] == "pdf_page_cleanup" and phase["status"] == "PASS" for phase in trace["phases"])

    def test_pdf_page_images_are_cleaned_after_ocr_failure(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        source_pdf = tmp_path / "inspection.pdf"
        source_pdf.write_bytes(b"%PDF-1.4")
        monkeypatch.setattr("pipeline.inspect_report._pdf_to_images", self._fake_pdf_pages)
        monkeypatch.setattr(
            "pipeline.inspect_report._ocr_images",
            lambda images, backend: (_ for _ in ()).throw(WorkbenchError("OCR failed")),
        )

        out_dir = tmp_path / "run"
        result = run_inspect_report(input_path=source_pdf, output_dir=out_dir, skip_docx=True)

        assert source_pdf.exists()
        assert not (out_dir / "pages").exists()
        assert any(phase["phase"] == "pdf_page_cleanup" and phase["status"] == "PASS" for phase in result["phases"])

    def test_image_pipeline_produces_all_artifacts(
        self, sample_image: Path, tmp_path: Path, fake_ocr_backend: PaddleOCRBackend,
    ) -> None:
        out_dir = tmp_path / "run_output"
        result = run_inspect_report(
            input_path=sample_image,
            output_dir=out_dir,
            ocr_backend=fake_ocr_backend,
        )
        assert result["status"] == "complete"
        assert (out_dir / "metadata.json").exists()
        assert (out_dir / "trace.json").exists()
        assert (out_dir / "ocr.txt").exists()
        assert (out_dir / "findings.json").exists()
        assert (out_dir / "network_report.json").exists()

        # Verify OCR text has page markers
        ocr_text = (out_dir / "ocr.txt").read_text()
        assert "--- PAGE 1 ---" in ocr_text

        # Verify findings JSON schema
        findings = json.loads((out_dir / "findings.json").read_text())
        assert "report_metadata" in findings
        assert "findings" in findings
        assert "uncertainties" in findings

        # Verify network report
        net = json.loads((out_dir / "network_report.json").read_text())
        assert net["external_connection_count"] == 0

    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    def test_image_pipeline_produces_docx(
        self, sample_image: Path, tmp_path: Path, fake_ocr_backend: PaddleOCRBackend,
    ) -> None:
        out_dir = tmp_path / "run_docx"
        result = run_inspect_report(
            input_path=sample_image,
            output_dir=out_dir,
            ocr_backend=fake_ocr_backend,
        )
        assert (out_dir / "approval_note.docx").exists()
        assert (out_dir / "approval_note.docx").stat().st_size > 0

    def test_invalid_input_fails_gracefully(self, tmp_path: Path) -> None:
        result = run_inspect_report(
            input_path=tmp_path / "nonexistent.pdf",
            output_dir=tmp_path / "run_fail",
        )
        assert result["status"] == "error"
        assert len(result["errors"]) > 0


# ---------------------------------------------------------------------------
# 9. CLI command/help behavior
# ---------------------------------------------------------------------------

class TestCLI:

    def test_cli_help_includes_inspect_report(self) -> None:
        result = subprocess.run(
            [sys.executable, str(_ROOT / "cli.py"), "--help"],
            capture_output=True, text=True, timeout=10,
        )
        assert result.returncode == 0
        assert "inspect-report" in result.stdout

    def test_inspect_report_help(self) -> None:
        result = subprocess.run(
            [sys.executable, str(_ROOT / "cli.py"), "inspect-report", "--help"],
            capture_output=True, text=True, timeout=10,
        )
        assert result.returncode == 0
        assert "--input" in result.stdout

    def test_inspect_report_missing_input(self) -> None:
        result = subprocess.run(
            [sys.executable, str(_ROOT / "cli.py"), "inspect-report"],
            capture_output=True, text=True, timeout=10,
        )
        assert result.returncode != 0
