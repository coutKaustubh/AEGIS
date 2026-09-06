from pathlib import Path

from tools.documents import DocumentTools


def test_document_discovery_and_bounded_search(tmp_path: Path):
    (tmp_path / "report.md").write_text("# Findings\n\nPump P-101 has high vibration.\n", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("private", encoding="utf-8")
    tools = DocumentTools(tmp_path)
    listed = tools.list_documents()
    assert listed["success"] is True
    assert {item["path"] for item in listed["documents"]} == {"report.md", "secret.txt"}
    result = tools.search_documents("P-101 vibration")
    assert result["success"] is True
    assert result["sections"][0]["source_path"] == "report.md"
    assert "P-101" in result["sections"][0]["excerpt"]


def test_document_paths_are_workspace_confined(tmp_path: Path):
    tools = DocumentTools(tmp_path)
    assert tools.extract_document_text("../outside.txt")["error"] == "OutsideWorkspace"
    (tmp_path / "unsupported.bin").write_bytes(b"x")
    assert tools.extract_document_text("unsupported.bin")["error"] == "UnsupportedFileType"


def test_docx_structure_is_returned_when_available(tmp_path: Path):
    from docx import Document
    source = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("Inspection", level=1)
    doc.add_paragraph("Valve V-204 requires review.")
    doc.save(source)
    result = DocumentTools(tmp_path).extract_document_text("report.docx")
    assert result["success"] is True
    assert any("V-204" in section["excerpt"] for section in result["sections"])

