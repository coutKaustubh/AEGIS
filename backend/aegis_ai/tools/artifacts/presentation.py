"""Deterministic PowerPoint specification, creation, editing, and validation."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from .common import ArtifactError, ArtifactResult, file_digest, json_spec, output_path, resolve_workspace_path


THEMES = {
    "simple": {"font": "Aptos", "accent": "1F4E79", "background": "FFFFFF"},
    "professional": {"font": "Aptos", "accent": "1F4E79", "background": "FFFFFF"},
    "technical": {"font": "Aptos", "accent": "0B6E8E", "background": "F7FAFC"},
    "business": {"font": "Aptos", "accent": "1F4E79", "background": "FFFFFF"},
    "academic": {"font": "Aptos", "accent": "5B2C83", "background": "FFFFFF"},
    "dark": {"font": "Aptos", "accent": "5CC8FF", "background": "17212B"},
    "light": {"font": "Aptos", "accent": "2F75B5", "background": "FFFFFF"},
}


def _color(value: str):
    from pptx.dml.color import RGBColor
    value = str(value).replace("#", "")
    return RGBColor.from_string((value or "1F4E79")[:6].upper())


def _add_text(slide, text: str, left: float, top: float, width: float, height: float, *, size: int = 20, bold: bool = False, color: str = "1F4E79"):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _color(color)
    return box


def _add_footer(slide, number: int, accent: str):
    from pptx.util import Inches
    _add_text(slide, str(number), 12.75, 7.12, 0.3, 0.2, size=9, color=accent)


def validate_presentation(path: str | Path, *, workspace_root: str | Path | None = None, expected: dict[str, Any] | None = None) -> dict[str, Any]:
    target = resolve_workspace_path(path, workspace_root or Path(path).parent) if workspace_root else Path(path).resolve()
    result: dict[str, Any] = {"status": "failed", "artifact_type": "pptx", "slides": 0, "warnings": [], "errors": []}
    if not target.is_file() or target.stat().st_size == 0:
        result["errors"].append("file_missing_or_empty")
        return result
    if not zipfile.is_zipfile(target):
        result["errors"].append("invalid_pptx_package")
        return result
    try:
        from pptx import Presentation
        presentation = Presentation(str(target))
        result["slides"] = len(presentation.slides)
        titles = []
        empty = 0
        for slide in presentation.slides:
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            titles.append(texts[0] if texts else "")
            empty += not bool(texts)
        result["titles"] = titles
        if empty:
            result["warnings"].append(f"{empty} empty slide(s)")
        expected = expected or {}
        if expected.get("slides") is not None and result["slides"] != int(expected["slides"]):
            result["errors"].append(f"expected_{expected['slides']}_slides_found_{result['slides']}")
        for text in expected.get("contains", []):
            if not any(str(text).casefold() in title.casefold() for title in titles):
                result["errors"].append(f"missing_expected_text:{text}")
        result["status"] = "passed" if not result["errors"] else "failed"
        result["sha256"] = file_digest(target)
    except Exception as exc:
        result["errors"].append(f"parse_error:{type(exc).__name__}")
    return result


def create_presentation(spec: dict[str, Any] | str, output: str | Path | None, workspace_root: str | Path, *, progress=None) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    data = json_spec(spec)
    theme_name = str(data.get("theme", "professional")).lower()
    theme = THEMES.get(theme_name, THEMES["professional"])
    title = str(data.get("title") or "AEGIS Presentation").strip()
    if "slides" in data and not isinstance(data.get("slides"), list):
        raise ArtifactError("invalid_specification", "Presentation slides must be a list")
    if "slides" in data and not data.get("slides"):
        raise ArtifactError("invalid_specification", "Presentation requires at least one slide")
    slides = data.get("slides") or [{"layout": "content", "title": title, "bullets": ["Sovereign local execution", "Deterministic artifact tools", "Validated workspace outputs"]}]
    if not isinstance(slides, list) or not slides:
        raise ArtifactError("invalid_specification", "Presentation requires at least one slide")
    path = output_path(output, workspace_root, artifact_type="pptx", stem="presentation")
    if progress: progress("artifact_generation_started", artifact_type="pptx", path=str(path))
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    # Remove the default empty slide when using a fresh presentation.
    while presentation.slides:
        relationship = presentation.slides._sldIdLst[0].rId
        presentation.part.drop_rel(relationship)
        del presentation.slides._sldIdLst[0]
    for index, raw in enumerate(slides, 1):
        slide_data = raw if isinstance(raw, dict) else {"title": str(raw), "bullets": []}
        layout = str(slide_data.get("layout", "title" if index == 1 else "content")).lower()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _color(theme["background"])
        if layout == "title" or index == 1:
            _add_text(slide, slide_data.get("title", title), 0.8, 2.0, 11.8, 1.0, size=36, bold=True, color=theme["accent"])
            _add_text(slide, slide_data.get("subtitle", data.get("subtitle", "")), 0.85, 3.15, 11.2, 0.7, size=20, color="52606D")
        else:
            _add_text(slide, slide_data.get("title", f"Slide {index}"), 0.65, 0.45, 12.0, 0.55, size=28, bold=True, color=theme["accent"])
            bullets = slide_data.get("bullets") or slide_data.get("content") or []
            if isinstance(bullets, str): bullets = [bullets]
            if layout == "two_column":
                columns = slide_data.get("columns") or [bullets, []]
                for col, items in enumerate(columns[:2]):
                    for row, item in enumerate(items if isinstance(items, list) else [items]):
                        _add_text(slide, f"• {item}", 0.8 + col * 6.15, 1.45 + row * 0.48, 5.6, 0.4, size=18, color="263238")
            else:
                for row, item in enumerate(bullets):
                    _add_text(slide, f"• {item}", 0.9, 1.45 + row * 0.55, 11.5, 0.45, size=19, color="263238")
            if slide_data.get("table"):
                table_data = slide_data["table"]
                rows = table_data.get("rows", [])
                cols = len(table_data.get("headers", [])) or (len(rows[0]) if rows else 1)
                table = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.8), Inches(4.7), Inches(11.8), Inches(1.6)).table
                for c, value in enumerate(table_data.get("headers", [])): table.cell(0, c).text = str(value)
                for r, values in enumerate(rows, 1):
                    for c, value in enumerate(values[:cols]): table.cell(r, c).text = str(value)
            chart = slide_data.get("chart")
            if chart:
                chart_data = CategoryChartData(); chart_data.categories = chart.get("categories", [])
                for series in chart.get("series", []): chart_data.add_series(str(series.get("name", "Series")), series.get("values", []))
                chart_type = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}.get(str(chart.get("type", "column")).lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)
                slide.shapes.add_chart(chart_type, Inches(7.0), Inches(1.35), Inches(5.5), Inches(3.0), chart_data)
        _add_footer(slide, index, theme["accent"])
    presentation.core_properties.title = title
    presentation.save(str(path))
    validation = validate_presentation(path, workspace_root=workspace_root, expected={"slides": len(slides), "contains": [title] if len(slides) == 1 else []})
    result = ArtifactResult("success" if validation["status"] == "passed" else "failure", "pptx", str(path), path.stat().st_size, validation, {"slides": len(slides), "theme": theme_name, "sha256": file_digest(path)}, validation.get("errors", [])).to_dict()
    if progress: progress("artifact_validation_passed" if result["status"] == "success" else "artifact_validation_failed", artifact_type="pptx", validation=validation)
    return result


def edit_presentation(source: str | Path, operations: list[dict[str, Any]], output: str | Path | None, workspace_root: str | Path, *, progress=None) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Inches
    source_path = resolve_workspace_path(source, workspace_root, must_exist=True)
    target = output_path(output, workspace_root, artifact_type="pptx", stem=f"edited_{source_path.stem}")
    presentation = Presentation(str(source_path))
    for operation in operations or []:
        name = str(operation.get("operation", "")).lower()
        index = int(operation.get("slide", operation.get("slide_index", 0)))
        if name in {"update_text", "replace_text"}:
            if index < 0 or index >= len(presentation.slides): raise ArtifactError("invalid_input", "Slide index is out of range")
            old = str(operation.get("old_text", "")); new = str(operation.get("new_text", operation.get("text", "")))
            changed = False
            for shape in presentation.slides[index].shapes:
                if hasattr(shape, "text") and (not old or old in shape.text):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs: run.text = run.text.replace(old, new)
                    changed = True
            if not changed: raise ArtifactError("invalid_input", "Requested presentation text was not found")
        elif name == "add_slide":
            slide = presentation.slides.add_slide(presentation.slide_layouts[6]); _add_text(slide, operation.get("title", "New slide"), 0.8, 0.6, 11, 0.7, size=28, bold=True)
        elif name == "remove_slide":
            if index < 0 or index >= len(presentation.slides): raise ArtifactError("invalid_input", "Slide index is out of range")
            relationship = presentation.slides._sldIdLst[index].rId; presentation.part.drop_rel(relationship); del presentation.slides._sldIdLst[index]
        else:
            raise ArtifactError("unsupported_operation", f"Unsupported presentation operation: {name}")
    presentation.save(str(target))
    validation = validate_presentation(target, workspace_root=workspace_root)
    return ArtifactResult("success" if validation["status"] == "passed" else "failure", "pptx", str(target), target.stat().st_size, validation, {"source": str(source_path), "operations": len(operations or [])}, validation.get("errors", [])).to_dict()
