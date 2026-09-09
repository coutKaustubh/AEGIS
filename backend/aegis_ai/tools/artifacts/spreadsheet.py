"""Deterministic Excel specification, creation, editing, and validation."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from .common import ArtifactError, ArtifactResult, file_digest, json_spec, output_path, resolve_workspace_path


TEMPLATES = {"basic_table", "financial", "budget", "analytics", "project_tracker", "data_analysis"}


def validate_workbook(path: str | Path, *, workspace_root: str | Path | None = None, expected: dict[str, Any] | None = None) -> dict[str, Any]:
    target = resolve_workspace_path(path, workspace_root or Path(path).parent) if workspace_root else Path(path).resolve()
    result: dict[str, Any] = {"status": "failed", "artifact_type": "xlsx", "sheets": [], "warnings": [], "errors": []}
    if not target.is_file() or target.stat().st_size == 0:
        result["errors"].append("file_missing_or_empty"); return result
    if not zipfile.is_zipfile(target):
        result["errors"].append("invalid_xlsx_package"); return result
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(str(target), data_only=False, read_only=False)
        result["sheets"] = list(workbook.sheetnames)
        expected = expected or {}
        for name in expected.get("sheets", []):
            if name not in workbook.sheetnames: result["errors"].append(f"missing_sheet:{name}")
        for sheet_name, cells in expected.get("cells", {}).items():
            if sheet_name not in workbook.sheetnames: continue
            sheet = workbook[sheet_name]
            for cell in cells:
                if sheet[cell].value is None: result["errors"].append(f"empty_expected_cell:{sheet_name}!{cell}")
        result["charts"] = sum(len(sheet._charts) for sheet in workbook.worksheets)
        result["tables"] = sum(len(sheet.tables) for sheet in workbook.worksheets)
        result["status"] = "passed" if not result["errors"] else "failed"
        result["sha256"] = file_digest(target)
    except Exception as exc:
        result["errors"].append(f"parse_error:{type(exc).__name__}")
    return result


def _style_sheet(sheet, spec: dict[str, Any]):
    from openpyxl.styles import Alignment, Font, PatternFill
    header_fill = PatternFill("solid", fgColor=str(spec.get("header_color", "1F4E79")).replace("#", "")[:6])
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = spec.get("freeze_panes", "A2")
    for index, width in enumerate(spec.get("column_widths", []), 1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else f"A{index}"].width = float(width)


def create_workbook(spec: dict[str, Any] | str, output: str | Path | None, workspace_root: str | Path, *, progress=None) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo

    data = json_spec(spec)
    sheets = data.get("worksheets") or data.get("sheets") or []
    title = str(data.get("title") or "AEGIS Workbook")
    if not sheets:
        sheets = [{"name": "Expenses", "headers": ["Month", "Category", "Amount"], "rows": [["January", "Operations", 0]], "formulas": {"D2": "=C2"}, "charts": [{"type": "column", "title": "Expenses", "data_range": "A1:C2"}]}]
    if not isinstance(sheets, list) or not sheets: raise ArtifactError("invalid_specification", "Workbook requires at least one worksheet")
    path = output_path(output, workspace_root, artifact_type="xlsx", stem="workbook")
    if progress: progress("artifact_generation_started", artifact_type="xlsx", path=str(path))
    workbook = Workbook(); workbook.remove(workbook.active)
    for sheet_spec in sheets:
        if not isinstance(sheet_spec, dict): raise ArtifactError("invalid_specification", "Worksheet specification must be an object")
        name = str(sheet_spec.get("name") or "Sheet")[:31]
        if name in workbook.sheetnames: raise ArtifactError("invalid_specification", f"Duplicate worksheet: {name}")
        sheet = workbook.create_sheet(name)
        headers = sheet_spec.get("headers") or []
        rows = sheet_spec.get("rows") or []
        if headers: sheet.append(list(headers))
        for row in rows: sheet.append(list(row) if isinstance(row, (list, tuple)) else [row])
        for cell, formula in (sheet_spec.get("formulas") or {}).items(): sheet[cell] = str(formula)
        for cell, value in (sheet_spec.get("cells") or {}).items(): sheet[cell] = value
        for merge in sheet_spec.get("merged_cells", []): sheet.merge_cells(str(merge))
        _style_sheet(sheet, sheet_spec)
        if sheet_spec.get("table") and sheet.max_row >= 2 and sheet.max_column >= 1:
            table_spec = sheet_spec["table"] if isinstance(sheet_spec["table"], dict) else {}
            ref = str(table_spec.get("ref") or f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
            existing_tables = sum(len(existing.tables) for existing in workbook.worksheets)
            table = Table(displayName=str(table_spec.get("name", f"Table_{existing_tables + 1}"))[:255], ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        for chart_spec in sheet_spec.get("charts", []):
            chart_cls = {"line": LineChart, "pie": PieChart, "bar": BarChart, "column": BarChart}.get(str(chart_spec.get("type", "column")).lower(), BarChart)
            chart = chart_cls(); chart.title = str(chart_spec.get("title", "Chart")); chart.style = 10
            reference = str(chart_spec.get("data_range", f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"))
            from openpyxl.utils.cell import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(reference.split("!", 1)[-1])
            data_ref = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data_ref, titles_from_data=True)
            if chart_cls is not PieChart: chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
            sheet.add_chart(chart, str(chart_spec.get("anchor", "F2")))
        for cell, fmt in (sheet_spec.get("number_formats") or {}).items(): sheet[cell].number_format = str(fmt)
    workbook.properties.title = title
    workbook.save(str(path))
    expected = {"sheets": [str(item.get("name"))[:31] for item in sheets if isinstance(item, dict)]}
    validation = validate_workbook(path, workspace_root=workspace_root, expected=expected)
    result = ArtifactResult("success" if validation["status"] == "passed" else "failure", "xlsx", str(path), path.stat().st_size, validation, {"sheets": validation.get("sheets", []), "title": title, "sha256": file_digest(path)}, validation.get("errors", [])).to_dict()
    if progress: progress("artifact_validation_passed" if result["status"] == "success" else "artifact_validation_failed", artifact_type="xlsx", validation=validation)
    return result


def edit_workbook(source: str | Path, operations: list[dict[str, Any]], output: str | Path | None, workspace_root: str | Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference
    source_path = resolve_workspace_path(source, workspace_root, must_exist=True)
    target = output_path(output, workspace_root, artifact_type="xlsx", stem=f"edited_{source_path.stem}")
    workbook = load_workbook(str(source_path))
    for operation in operations or []:
        name = str(operation.get("operation", "")).lower(); sheet_name = str(operation.get("sheet", ""))
        if name == "create_sheet": workbook.create_sheet(sheet_name or "Sheet")
        elif name == "rename_sheet":
            if sheet_name not in workbook.sheetnames: raise ArtifactError("invalid_input", f"Unknown worksheet: {sheet_name}")
            workbook[sheet_name].title = str(operation.get("new_name", "Sheet"))[:31]
        elif name == "delete_sheet":
            if sheet_name not in workbook.sheetnames: raise ArtifactError("invalid_input", f"Unknown worksheet: {sheet_name}")
            if len(workbook.sheetnames) == 1: raise ArtifactError("invalid_operation", "Workbook must retain one worksheet")
            del workbook[sheet_name]
        elif name in {"write_range", "write_cell"}:
            if sheet_name not in workbook.sheetnames: raise ArtifactError("invalid_input", f"Unknown worksheet: {sheet_name}")
            sheet = workbook[sheet_name]; start = str(operation.get("cell", operation.get("start", "A1"))); values = operation.get("values", operation.get("value"))
            if name == "write_cell": sheet[start] = values
            else:
                for row_index, row in enumerate(values or [], sheet[start].row):
                    for col_index, value in enumerate(row if isinstance(row, list) else [row], sheet[start].column): sheet.cell(row_index, col_index).value = value
        elif name == "add_formula":
            if sheet_name not in workbook.sheetnames: raise ArtifactError("invalid_input", f"Unknown worksheet: {sheet_name}")
            workbook[sheet_name][str(operation.get("cell", "A1"))] = str(operation.get("formula", ""))
        elif name == "freeze_panes": workbook[sheet_name].freeze_panes = str(operation.get("cell", "A2"))
        elif name == "add_chart":
            sheet = workbook[sheet_name]; chart = BarChart(); chart.title = str(operation.get("title", "Chart"))
            from openpyxl.utils.cell import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(str(operation.get("data_range", "A1:B2")).split("!", 1)[-1])
            chart.add_data(Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row), titles_from_data=True); sheet.add_chart(chart, str(operation.get("anchor", "F2")))
        else: raise ArtifactError("unsupported_operation", f"Unsupported spreadsheet operation: {name}")
    workbook.save(str(target)); validation = validate_workbook(target, workspace_root=workspace_root)
    return ArtifactResult("success" if validation["status"] == "passed" else "failure", "xlsx", str(target), target.stat().st_size, validation, {"source": str(source_path), "operations": len(operations or [])}, validation.get("errors", [])).to_dict()
