"""Workspace-confined deterministic document discovery and lexical retrieval.

This is intentionally not a vector database or RAG system. It extracts small,
source-labelled sections so the existing Master/agent graph can reason over
bounded evidence without loading whole documents into context.
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from langchain_core.tools import BaseTool, tool


SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".pptx"}
MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_EXCERPT_CHARS = 1200


class DocumentTools:
    def __init__(self, root: str | Path, *, max_bytes: int = MAX_DOCUMENT_BYTES) -> None:
        self.root = Path(root).resolve()
        self.max_bytes = max_bytes

    def _resolve(self, raw: str | Path) -> Path | dict[str, Any]:
        candidate = (self.root / raw).resolve() if not Path(raw).is_absolute() else Path(raw).resolve()
        try:
            candidate.relative_to(self.root)
        except ValueError:
            return {"success": False, "error": "OutsideWorkspace", "message": "Document path is outside the approved workspace."}
        if not candidate.is_file():
            return {"success": False, "error": "FileNotFound", "message": "Document does not exist."}
        if candidate.stat().st_size > self.max_bytes:
            return {"success": False, "error": "DocumentTooLarge", "message": f"Document exceeds {self.max_bytes} bytes."}
        if candidate.suffix.lower() not in SUPPORTED_EXTENSIONS:
            return {"success": False, "error": "UnsupportedFileType", "message": f"Unsupported document type: {candidate.suffix}"}
        return candidate

    def list_documents(self, path: str = ".") -> dict[str, Any]:
        base = (self.root / path).resolve() if not Path(path).is_absolute() else Path(path).resolve()
        try:
            base.relative_to(self.root)
        except ValueError:
            return {"success": False, "error": "OutsideWorkspace", "documents": []}
        if not base.exists() or not base.is_dir():
            return {"success": False, "error": "DirectoryNotFound", "documents": []}
        documents = []
        for item in sorted(base.rglob("*")):
            if item.is_file() and item.suffix.lower() in SUPPORTED_EXTENSIONS and not any(part.startswith(".") for part in item.relative_to(self.root).parts):
                stat = item.stat()
                documents.append({"document_id": hashlib.sha256(str(item).encode()).hexdigest()[:16],
                                  "path": str(item.relative_to(self.root)), "format": item.suffix.lower().lstrip("."),
                                  "size_bytes": stat.st_size})
        return {"success": True, "documents": documents, "truncated": False}

    def inspect_document_metadata(self, path: str) -> dict[str, Any]:
        resolved = self._resolve(path)
        if isinstance(resolved, dict):
            return resolved
        stat = resolved.stat()
        result = {"success": True, "document_id": hashlib.sha256(str(resolved).encode()).hexdigest()[:16],
                  "source_path": str(resolved.relative_to(self.root)), "format": resolved.suffix.lower().lstrip("."),
                  "size_bytes": stat.st_size, "modified_at": stat.st_mtime}
        if resolved.suffix.lower() == ".pdf":
            try:
                import pymupdf
                doc = pymupdf.open(str(resolved))
                result["page_count"] = doc.page_count
                doc.close()
            except Exception as exc:
                result["metadata_error"] = type(exc).__name__
        return result

    def _sections(self, path: Path) -> list[dict[str, Any]]:
        suffix = path.suffix.lower()
        sections: list[dict[str, Any]] = []
        if suffix in {".txt", ".md", ".markdown"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            offset = 0
            for index, block in enumerate(re.split(r"\n\s*\n", text)):
                if block.strip():
                    start = text.find(block, offset)
                    sections.append({"section": f"paragraph-{index + 1}", "page": None,
                                     "excerpt": block[:MAX_EXCERPT_CHARS], "start_offset": start,
                                     "end_offset": start + len(block), "source_path": str(path.relative_to(self.root))})
                    offset = start + len(block)
        elif suffix == ".pdf":
            import pymupdf
            doc = pymupdf.open(str(path))
            for page_number, page in enumerate(doc, 1):
                text = page.get_text("text") or ""
                if text.strip():
                    sections.append({"section": f"page-{page_number}", "page": page_number,
                                     "excerpt": text[:MAX_EXCERPT_CHARS], "start_offset": 0,
                                     "end_offset": len(text), "source_path": str(path.relative_to(self.root))})
            doc.close()
        elif suffix == ".docx":
            from docx import Document
            doc = Document(str(path))
            for index, paragraph in enumerate(doc.paragraphs):
                if paragraph.text.strip():
                    sections.append({"section": paragraph.style.name or f"paragraph-{index + 1}", "page": None,
                                     "excerpt": paragraph.text[:MAX_EXCERPT_CHARS], "start_offset": None,
                                     "end_offset": None, "source_path": str(path.relative_to(self.root))})
            for table_index, table in enumerate(doc.tables):
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                text = "\n".join(rows)
                if text.strip():
                    sections.append({"section": f"table-{table_index + 1}", "page": None,
                                     "excerpt": text[:MAX_EXCERPT_CHARS], "start_offset": None,
                                     "end_offset": None, "source_path": str(path.relative_to(self.root))})
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(path))
            for slide_number, slide in enumerate(presentation.slides, 1):
                text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
                if text.strip():
                    sections.append({"section": f"slide-{slide_number}", "page": slide_number,
                                     "excerpt": text[:MAX_EXCERPT_CHARS], "start_offset": None,
                                     "end_offset": None, "source_path": str(path.relative_to(self.root))})
        return sections

    def extract_document_text(self, path: str, max_sections: int = 100) -> dict[str, Any]:
        resolved = self._resolve(path)
        if isinstance(resolved, dict):
            return resolved
        try:
            sections = self._sections(resolved)
            return {"success": True, "document_id": hashlib.sha256(str(resolved).encode()).hexdigest()[:16],
                    "source_path": str(resolved.relative_to(self.root)), "sections": sections[:max(1, min(max_sections, 100))],
                    "truncated": len(sections) > max_sections, "error": None}
        except Exception as exc:
            return {"success": False, "source_path": str(resolved.relative_to(self.root)),
                    "error": type(exc).__name__, "message": str(exc)[:300]}

    def search_documents(self, query: str, path: str = ".", max_results: int = 20) -> dict[str, Any]:
        if not query or len(query) > 500:
            return {"success": False, "error": "InvalidQuery", "sections": []}
        listed = self.list_documents(path)
        if not listed.get("success"):
            return listed | {"sections": []}
        terms = [term for term in re.findall(r"\w+", query.casefold()) if term]
        matches: list[dict[str, Any]] = []
        for item in listed["documents"]:
            extracted = self.extract_document_text(item["path"], max_sections=100)
            for section in extracted.get("sections", []):
                haystack = section.get("excerpt", "").casefold()
                if all(term in haystack for term in terms):
                    matches.append({**section, "document_id": item["document_id"], "query": query})
                    if len(matches) >= max(1, min(max_results, 50)):
                        return {"success": True, "query": query, "sections": matches, "truncated": True}
        return {"success": True, "query": query, "sections": matches, "truncated": False}

    def read_document_section(self, path: str, section: str = "", page: int | None = None) -> dict[str, Any]:
        extracted = self.extract_document_text(path)
        if not extracted.get("success"):
            return extracted
        sections = [item for item in extracted["sections"] if (not section or item.get("section") == section) and (page is None or item.get("page") == page)]
        return {**extracted, "sections": sections[:20], "truncated": len(sections) > 20}

    def as_langchain_tools(self) -> list[BaseTool]:
        @tool("list_documents")
        def list_documents(path: str = ".") -> dict[str, Any]:
            """List supported documents inside the approved workspace."""
            return self.list_documents(path)

        @tool("inspect_document_metadata")
        def inspect_document_metadata(path: str) -> dict[str, Any]:
            """Inspect bounded metadata for a workspace document."""
            return self.inspect_document_metadata(path)

        @tool("extract_document_text")
        def extract_document_text(path: str, max_sections: int = 100) -> dict[str, Any]:
            """Extract source-labelled sections from a local document."""
            return self.extract_document_text(path, max_sections)

        @tool("search_documents")
        def search_documents(query: str, path: str = ".", max_results: int = 20) -> dict[str, Any]:
            """Search supported workspace documents using bounded lexical matching."""
            return self.search_documents(query, path, max_results)

        @tool("read_document_section")
        def read_document_section(path: str, section: str = "", page: int | None = None) -> dict[str, Any]:
            """Read selected document sections without returning whole files."""
            return self.read_document_section(path, section, page)

        return [list_documents, inspect_document_metadata, extract_document_text, search_documents, read_document_section]

